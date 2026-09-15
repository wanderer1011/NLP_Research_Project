"""Generate a polished presentation deck for the Similarity Pipeline."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# Light theme palette aligned with iap_pipeline/create_presentation.py
BG_LIGHT = RGBColor(0xF5, 0xF6, 0xFA)
BG_CARD = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x00, 0x78, 0xD4)
ACCENT2 = RGBColor(0x6B, 0x3F, 0xA0)
TEXT_DARK = RGBColor(0x1E, 0x1E, 0x2E)
TEXT_MED = RGBColor(0x55, 0x55, 0x66)
GREEN = RGBColor(0x0E, 0x8A, 0x4E)
ORANGE = RGBColor(0xD8, 0x7B, 0x00)
RED = RGBColor(0xC5, 0x0F, 0x1F)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create a similarity pipeline PPTX report.")
    parser.add_argument(
        "--results-dir",
        default="con_emotion_linear",
        help="Directory containing metrics.json and predictions.csv (default: con_emotion_linear)",
    )
    parser.add_argument(
        "--output",
        default="Similarity_Pipeline_Presentation.pptx",
        help="Output PPTX path (default: Similarity_Pipeline_Presentation.pptx)",
    )
    parser.add_argument(
        "--title",
        default="Similarity / Retrieval-Based Manipulation Detection",
        help="Deck title shown on slide 1",
    )
    return parser.parse_args()


def set_slide_bg(slide, color: RGBColor) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, fill_color: RGBColor, border_color: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    font_size: int = 18,
    color: RGBColor = TEXT_DARK,
    bold: bool = False,
    alignment=PP_ALIGN.LEFT,
    font_name: str = "Segoe UI",
):
    tx_box = slide.shapes.add_textbox(left, top, width, height)
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tx_box


def add_bullet_text(
    slide,
    left,
    top,
    width,
    height,
    items: list[str],
    font_size: int = 16,
    color: RGBColor = TEXT_MED,
    bullet_color: RGBColor = ACCENT,
):
    tx_box = slide.shapes.add_textbox(left, top, width, height)
    tf = tx_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(6)
        p.space_after = Pt(4)

        run_bullet = p.add_run()
        run_bullet.text = "▸ "
        run_bullet.font.size = Pt(font_size)
        run_bullet.font.color.rgb = bullet_color
        run_bullet.font.name = "Segoe UI"

        run_text = p.add_run()
        run_text.text = item
        run_text.font.size = Pt(font_size)
        run_text.font.color.rgb = color
        run_text.font.name = "Segoe UI"
    return tx_box


def _read_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def _pct(value: float | int | None) -> str:
    if value is None:
        return "N/A"
    return f"{float(value) * 100:.1f}%"


def _float(value: object, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _clip_text(text: str, max_len: int = 240) -> str:
    compact = " ".join((text or "").split())
    if len(compact) <= max_len:
        return compact
    return compact[: max_len - 3] + "..."


def load_metrics(results_dir: Path) -> dict[str, float | int | str | bool | None]:
    metrics_path = results_dir / "metrics.json"
    aggregate_path = results_dir / "aggregate_metrics.json"

    if aggregate_path.exists():
        payload = _read_json(aggregate_path)
        mean = payload.get("mean", {}) if isinstance(payload, dict) else {}
        run_count = payload.get("num_runs", 0) if isinstance(payload, dict) else 0
        return {
            "accuracy": _float(mean.get("accuracy")),
            "precision": _float(mean.get("precision")),
            "recall": _float(mean.get("recall")),
            "f1": _float(mean.get("f1")),
            "false_negative_rate": _float(mean.get("false_negative_rate")),
            "false_positive_rate": _float(mean.get("false_positive_rate")),
            "pr_auc": _float(mean.get("pr_auc")),
            "roc_auc": _float(mean.get("roc_auc")),
            "technique_micro_f1": _float(mean.get("technique_micro_f1")),
            "vulnerability_micro_f1": _float(mean.get("vulnerability_micro_f1")),
            "threshold": _float(mean.get("threshold"), 0.5),
            "active_retrieval_mode": "aggregate",
            "active_technique_prediction_mode": "aggregate",
            "num_runs": int(run_count),
            "tp": 0,
            "tn": 0,
            "fp": 0,
            "fn": 0,
        }

    if not metrics_path.exists():
        raise FileNotFoundError(f"Could not find metrics.json or aggregate_metrics.json in {results_dir}")

    payload = _read_json(metrics_path)
    return {
        "accuracy": _float(payload.get("accuracy")),
        "precision": _float(payload.get("precision")),
        "recall": _float(payload.get("recall")),
        "f1": _float(payload.get("f1")),
        "false_negative_rate": _float(payload.get("false_negative_rate")),
        "false_positive_rate": _float(payload.get("false_positive_rate")),
        "pr_auc": _float(payload.get("pr_auc")),
        "roc_auc": _float(payload.get("roc_auc")),
        "technique_micro_f1": _float(payload.get("technique_micro_f1")),
        "vulnerability_micro_f1": _float(payload.get("vulnerability_micro_f1")),
        "threshold": _float(payload.get("threshold"), 0.5),
        "active_retrieval_mode": str(payload.get("active_retrieval_mode", "unknown")),
        "active_technique_prediction_mode": str(payload.get("active_technique_prediction_mode", "unknown")),
        "num_runs": 1,
        "tp": int(payload.get("tp", 0) or 0),
        "tn": int(payload.get("tn", 0) or 0),
        "fp": int(payload.get("fp", 0) or 0),
        "fn": int(payload.get("fn", 0) or 0),
    }


def load_prediction_summary(results_dir: Path) -> dict[str, object]:
    pred_path = results_dir / "predictions.csv"
    if not pred_path.exists():
        return {
            "sample_text": "predictions.csv not found",
            "sample_gt": "N/A",
            "sample_pred": "N/A",
            "sample_score": "N/A",
            "top_predicted_techniques": [],
        }

    technique_counter: Counter[str] = Counter()
    rows: list[dict[str, str]] = []

    with pred_path.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(row)
            pred_tech = (row.get("predicted_techniques", "") or "").strip()
            if pred_tech:
                for label in pred_tech.split("|"):
                    label = label.strip()
                    if label:
                        technique_counter[label] += 1

    if not rows:
        return {
            "sample_text": "No prediction rows available",
            "sample_gt": "N/A",
            "sample_pred": "N/A",
            "sample_score": "N/A",
            "top_predicted_techniques": [],
        }

    # Prefer a true-positive example first for demonstration clarity.
    sample = None
    for row in rows:
        if row.get("ground_truth") == "1" and row.get("predicted") == "1":
            sample = row
            break
    if sample is None:
        sample = rows[0]

    return {
        "sample_text": _clip_text(sample.get("text", "") or ""),
        "sample_gt": sample.get("ground_truth", "N/A"),
        "sample_pred": sample.get("predicted", "N/A"),
        "sample_score": sample.get("score", "N/A"),
        "top_predicted_techniques": [label for label, _ in technique_counter.most_common(4)],
    }


def build_presentation(title: str, results_dir: Path, output_path: Path) -> None:
    metrics = load_metrics(results_dir)
    prediction_summary = load_prediction_summary(results_dir)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide1, BG_LIGHT)
    add_shape(slide1, Inches(0), Inches(0), Inches(13.333), Inches(0.08), ACCENT)

    add_text(
        slide1,
        Inches(1.1),
        Inches(1.8),
        Inches(11.2),
        Inches(1.2),
        title,
        font_size=40,
        color=TEXT_DARK,
        bold=True,
        alignment=PP_ALIGN.CENTER,
    )

    add_text(
        slide1,
        Inches(1.4),
        Inches(3.1),
        Inches(10.6),
        Inches(0.9),
        "Memory-centric classification with semantic retrieval, threshold tuning, and hybrid label prediction",
        font_size=21,
        color=ACCENT,
        alignment=PP_ALIGN.CENTER,
    )

    add_shape(slide1, Inches(4.5), Inches(4.2), Inches(4.3), Inches(0.04), ACCENT2)

    add_shape(slide1, Inches(2.1), Inches(4.8), Inches(9.0), Inches(1.4), BG_CARD, ACCENT2)
    add_text(
        slide1,
        Inches(2.4),
        Inches(4.95),
        Inches(8.4),
        Inches(1.1),
        f"Results source: {results_dir.name}  |  Retrieval mode: {metrics['active_retrieval_mode']}  |  "
        f"Technique mode: {metrics['active_technique_prediction_mode']}",
        font_size=16,
        color=TEXT_MED,
        alignment=PP_ALIGN.CENTER,
    )

    add_text(
        slide1,
        Inches(1),
        Inches(6.6),
        Inches(11.3),
        Inches(0.5),
        "Built with Python  ·  sentence-transformers  ·  scikit-learn  ·  MentalManip Dataset",
        font_size=13,
        color=TEXT_MED,
        alignment=PP_ALIGN.CENTER,
    )

    # Slide 2: Pipeline architecture
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide2, BG_LIGHT)

    add_text(slide2, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7), "Architecture & Execution Flow", font_size=34, color=TEXT_DARK, bold=True)
    add_shape(slide2, Inches(0.6), Inches(0.95), Inches(3), Inches(0.05), ACCENT)

    cards = [
        (
            "STEP 1",
            "Prepare Data",
            "data_loader.py",
            [
                "Load MentalManip dialogues + labels",
                "Normalize text for stable embeddings",
                "Split technique and vulnerability labels",
            ],
            Inches(0.4),
        ),
        (
            "STEP 2",
            "Encode Dialogues",
            "SentenceTransformer",
            [
                "Build emotion/semantic embeddings",
                "Optional embedding cache for faster reruns",
                "Support batch inference on large datasets",
            ],
            Inches(4.6),
        ),
        (
            "STEP 3",
            "Retrieve + Decide",
            "pipeline.py",
            [
                "Global/per-class/linear scoring modes",
                "Threshold tuning with recall/FPR constraints",
                "Hybrid technique/vulnerability label prediction",
            ],
            Inches(8.8),
        ),
    ]

    for step_label, section_title, filename, bullets, left in cards:
        add_shape(slide2, left, Inches(1.5), Inches(3.9), Inches(4.8), BG_CARD, ACCENT2)
        add_shape(slide2, left + Inches(0.2), Inches(1.7), Inches(1.4), Inches(0.45), ACCENT2)
        add_text(
            slide2,
            left + Inches(0.25),
            Inches(1.72),
            Inches(1.3),
            Inches(0.4),
            step_label,
            font_size=13,
            color=RGBColor(0xFF, 0xFF, 0xFF),
            bold=True,
            alignment=PP_ALIGN.CENTER,
        )
        add_text(slide2, left + Inches(0.3), Inches(2.3), Inches(3.3), Inches(0.5), section_title, font_size=22, color=ACCENT, bold=True)
        add_text(slide2, left + Inches(0.3), Inches(2.8), Inches(3.3), Inches(0.4), filename, font_size=13, color=ORANGE)
        add_bullet_text(slide2, left + Inches(0.3), Inches(3.3), Inches(3.4), Inches(2.8), bullets, font_size=14, color=TEXT_MED)

    for arrow_left in [Inches(4.35), Inches(8.55)]:
        add_text(slide2, arrow_left, Inches(3.5), Inches(0.4), Inches(0.5), "→", font_size=36, color=ACCENT, bold=True, alignment=PP_ALIGN.CENTER)

    add_shape(slide2, Inches(0.4), Inches(6.55), Inches(12.5), Inches(0.7), BG_CARD, RGBColor(0xDD, 0xDD, 0xEE))
    add_text(
        slide2,
        Inches(0.6),
        Inches(6.6),
        Inches(12),
        Inches(0.55),
        "Key outputs: predictions.csv, metrics.json, threshold_curve.json, aggregate_metrics.json (multi-run)",
        font_size=13,
        color=TEXT_MED,
        alignment=PP_ALIGN.CENTER,
    )

    # Slide 3: Technical decisions
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide3, BG_LIGHT)

    add_text(slide3, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7), "Key Technical Decisions", font_size=34, color=TEXT_DARK, bold=True)
    add_shape(slide3, Inches(0.6), Inches(0.95), Inches(3), Inches(0.05), ACCENT)

    add_shape(slide3, Inches(0.4), Inches(1.4), Inches(6.2), Inches(2.5), BG_CARD, ACCENT2)
    add_text(slide3, Inches(0.7), Inches(1.55), Inches(5.5), Inches(0.5), "Retrieval Strategy", font_size=20, color=ACCENT, bold=True)
    add_bullet_text(
        slide3,
        Inches(0.7),
        Inches(2.1),
        Inches(5.7),
        Inches(1.7),
        [
            "Supports auto mode across global, per-class, and linear scorers",
            "Similarity-power weighting emphasizes close exemplars",
            "Class-wise margin features improve interpretability",
        ],
        font_size=14,
    )

    add_shape(slide3, Inches(6.8), Inches(1.4), Inches(6.2), Inches(2.5), BG_CARD, ACCENT2)
    add_text(slide3, Inches(7.1), Inches(1.55), Inches(5.5), Inches(0.5), "Decision Calibration", font_size=20, color=ACCENT, bold=True)
    add_bullet_text(
        slide3,
        Inches(7.1),
        Inches(2.1),
        Inches(5.7),
        Inches(1.7),
        [
            "Threshold tuned on validation split",
            "Constraint-aware optimization: recall target + optional max FPR",
            "Optional Platt / isotonic probability calibration",
        ],
        font_size=14,
    )

    add_shape(slide3, Inches(0.4), Inches(4.2), Inches(6.2), Inches(2.5), BG_CARD, ACCENT2)
    add_text(slide3, Inches(0.7), Inches(4.35), Inches(5.5), Inches(0.5), "Label Attribution", font_size=20, color=ACCENT, bold=True)
    add_bullet_text(
        slide3,
        Inches(0.7),
        Inches(4.9),
        Inches(5.7),
        Inches(1.7),
        [
            "Neighbor vote model for technique/vulnerability labels",
            "One-vs-rest linear predictor for techniques",
            "Hybrid blending balances precision and recall",
        ],
        font_size=14,
    )

    add_shape(slide3, Inches(6.8), Inches(4.2), Inches(6.2), Inches(2.5), BG_CARD, ACCENT2)
    add_text(slide3, Inches(7.1), Inches(4.35), Inches(5.5), Inches(0.5), "Operational Efficiency", font_size=20, color=ACCENT, bold=True)
    add_bullet_text(
        slide3,
        Inches(7.1),
        Inches(4.9),
        Inches(5.7),
        Inches(1.7),
        [
            "No end-to-end deep classifier training loop",
            "Embedding cache for iterative sweeps",
            "Configurable multi-run reproducibility via seeds",
        ],
        font_size=14,
    )

    # Slide 4: Metrics snapshot
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide4, BG_LIGHT)

    add_text(slide4, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7), "Results Snapshot", font_size=34, color=TEXT_DARK, bold=True)
    add_shape(slide4, Inches(0.6), Inches(0.95), Inches(3), Inches(0.05), ACCENT)

    metric_cards = [
        (_pct(metrics["accuracy"]), "Accuracy", ACCENT),
        (_pct(metrics["precision"]), "Precision", ACCENT2),
        (_pct(metrics["recall"]), "Recall", GREEN),
        (_pct(metrics["f1"]), "F1", GREEN),
        (_pct(metrics["false_negative_rate"]), "FN Rate", ORANGE),
    ]

    for i, (value, label, clr) in enumerate(metric_cards):
        left = Inches(0.4 + i * 2.6)
        add_shape(slide4, left, Inches(1.4), Inches(2.35), Inches(2.0), BG_CARD, clr)
        add_text(slide4, left, Inches(1.65), Inches(2.35), Inches(0.8), value, font_size=36, color=clr, bold=True, alignment=PP_ALIGN.CENTER)
        add_text(slide4, left, Inches(2.45), Inches(2.35), Inches(0.5), label, font_size=16, color=TEXT_MED, alignment=PP_ALIGN.CENTER)

    add_shape(slide4, Inches(0.4), Inches(3.8), Inches(12.5), Inches(3.1), BG_CARD, ACCENT2)
    add_text(slide4, Inches(0.7), Inches(3.95), Inches(6.0), Inches(0.5), "Secondary Metrics", font_size=22, color=ACCENT, bold=True)
    add_bullet_text(
        slide4,
        Inches(0.7),
        Inches(4.5),
        Inches(6.0),
        Inches(2.2),
        [
            f"PR-AUC: {_pct(metrics['pr_auc'])}  |  ROC-AUC: {_pct(metrics['roc_auc'])}",
            f"Technique micro-F1: {_pct(metrics['technique_micro_f1'])}",
            f"Vulnerability micro-F1: {_pct(metrics['vulnerability_micro_f1'])}",
            f"Decision threshold: {metrics['threshold']:.3f}  |  Runs: {metrics['num_runs']}",
        ],
        font_size=14,
        color=TEXT_MED,
    )

    add_text(slide4, Inches(7.0), Inches(4.5), Inches(5.4), Inches(0.5), "Confusion Counts", font_size=16, color=ACCENT, bold=True)
    add_shape(slide4, Inches(7.0), Inches(5.0), Inches(5.5), Inches(1.6), RGBColor(0xF0, 0xF0, 0xF5))
    add_text(
        slide4,
        Inches(7.2),
        Inches(5.15),
        Inches(5.1),
        Inches(1.3),
        f"TP={metrics['tp']}   TN={metrics['tn']}\nFP={metrics['fp']}   FN={metrics['fn']}",
        font_size=20,
        color=TEXT_DARK,
        font_name="Consolas",
        alignment=PP_ALIGN.CENTER,
    )

    # Slide 5: Qualitative view
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide5, BG_LIGHT)

    add_text(slide5, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7), "Qualitative Prediction View", font_size=34, color=TEXT_DARK, bold=True)
    add_shape(slide5, Inches(0.6), Inches(0.95), Inches(3), Inches(0.05), ACCENT)

    top_techniques = prediction_summary["top_predicted_techniques"]
    top_techniques_text = ", ".join(top_techniques) if top_techniques else "No technique predictions available"

    add_shape(slide5, Inches(0.4), Inches(1.5), Inches(12.5), Inches(2.0), BG_CARD, ACCENT2)
    add_text(
        slide5,
        Inches(0.7),
        Inches(1.7),
        Inches(11.9),
        Inches(1.6),
        f"Sample dialogue: {prediction_summary['sample_text']}",
        font_size=14,
        color=TEXT_MED,
    )

    add_shape(slide5, Inches(0.4), Inches(3.8), Inches(6.1), Inches(2.9), BG_CARD, ACCENT2)
    add_text(slide5, Inches(0.7), Inches(4.0), Inches(5.5), Inches(0.5), "Sample Decision", font_size=22, color=ACCENT, bold=True)
    add_bullet_text(
        slide5,
        Inches(0.7),
        Inches(4.55),
        Inches(5.6),
        Inches(2.0),
        [
            f"Ground truth label: {prediction_summary['sample_gt']}",
            f"Predicted label: {prediction_summary['sample_pred']}",
            f"Model score: {prediction_summary['sample_score']}",
            f"Threshold: {metrics['threshold']:.3f}",
        ],
        font_size=14,
    )

    add_shape(slide5, Inches(6.8), Inches(3.8), Inches(6.1), Inches(2.9), BG_CARD, ACCENT2)
    add_text(slide5, Inches(7.1), Inches(4.0), Inches(5.5), Inches(0.5), "Most Frequent Predicted Techniques", font_size=22, color=ACCENT, bold=True)
    add_bullet_text(
        slide5,
        Inches(7.1),
        Inches(4.55),
        Inches(5.6),
        Inches(2.0),
        [
            top_techniques_text,
            "Useful for auditing what behaviors the retriever focuses on",
            "Can reveal overprediction patterns for specific tactics",
        ],
        font_size=14,
    )

    # Slide 6: Impact and limitations
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide6, BG_LIGHT)

    add_text(slide6, Inches(0.6), Inches(0.3), Inches(12), Inches(0.7), "Impact, Tradeoffs, Next Steps", font_size=34, color=TEXT_DARK, bold=True)
    add_shape(slide6, Inches(0.6), Inches(0.95), Inches(3), Inches(0.05), ORANGE)

    add_shape(slide6, Inches(0.4), Inches(1.5), Inches(6.2), Inches(4.9), BG_CARD, ACCENT2)
    add_text(slide6, Inches(0.7), Inches(1.7), Inches(5.6), Inches(0.5), "Strengths", font_size=22, color=ACCENT, bold=True)
    add_bullet_text(
        slide6,
        Inches(0.7),
        Inches(2.25),
        Inches(5.7),
        Inches(3.8),
        [
            "Interpretable via nearest-exemplar evidence",
            "Configurable objective allows recall-vs-FPR balancing",
            "Easy iterative improvement with cached embeddings and reruns",
            "Hybrid labeling improves behavioral signal coverage",
        ],
        font_size=14,
    )

    add_shape(slide6, Inches(6.8), Inches(1.5), Inches(6.2), Inches(4.9), BG_CARD, ORANGE)
    add_text(slide6, Inches(7.1), Inches(1.7), Inches(5.6), Inches(0.5), "Limitations / Risks", font_size=22, color=ORANGE, bold=True)
    add_bullet_text(
        slide6,
        Inches(7.1),
        Inches(2.25),
        Inches(5.7),
        Inches(3.8),
        [
            "Performance is bounded by embedding quality and dataset coverage",
            "Novel phrasing can reduce nearest-neighbor reliability",
            "Threshold transferability may drift across domains",
            "Label predictions depend on annotation completeness",
        ],
        font_size=14,
        bullet_color=ORANGE,
    )

    add_text(
        slide6,
        Inches(0.6),
        Inches(6.7),
        Inches(12),
        Inches(0.5),
        "Recommended next steps: compare modes across multiple seeds and audit high-confidence false positives.",
        font_size=12,
        color=TEXT_MED,
        alignment=PP_ALIGN.CENTER,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


if __name__ == "__main__":
    args = parse_args()

    base_dir = Path(__file__).resolve().parent
    results_dir = Path(args.results_dir)
    if not results_dir.is_absolute():
        results_dir = base_dir / results_dir

    output_path = Path(args.output)
    if not output_path.is_absolute():
        output_path = base_dir / output_path

    build_presentation(args.title, results_dir=results_dir, output_path=output_path)
    print(f"Presentation saved to: {output_path}")
